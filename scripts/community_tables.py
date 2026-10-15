"""Export Word table evidence locally. No model calls, uploads, or approved triples."""
import argparse
import hashlib
import json
from pathlib import Path
from docx import Document


def text(cell):
    return '\n'.join(p.text for p in cell.paragraphs).strip()


def extract_document(path):
    document = Document(path)
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    records, warnings = [], []
    # These layouts were inspected in the first document. Each other file must
    # be checked against its original; equal dimensions do not prove equal meaning.
    layouts = {1: (5, 2, 'resource_category'),
               2: (7, 1, 'month'), 3: (7, 1, 'disaster')}
    for table_index, table in enumerate(document.tables, 1):
        if table_index not in layouts:
            warnings.append(f'Table {table_index}: unconfigured layout; not exported')
            continue
        width, first_response, context_name = layouts[table_index]
        if not table.rows or any(len(r.cells) != width for r in table.rows):
            warnings.append(f'Table {table_index}: unexpected grid; not exported')
            continue
        headers = table.rows[0].cells
        for row_index, row in enumerate(table.rows[1:], 2):
            context = text(row.cells[0])
            seen = set()
            for column_index in range(first_response, width):
                cell = row.cells[column_index]
                # A vertically merged cell may intentionally apply to several
                # rows; preserve each row context. Horizontal duplicates in the
                # response area cannot safely be assigned to one header.
                if cell._tc in seen:
                    continue
                seen.add(cell._tc)
                value = text(cell)
                if not value:
                    continue  # Blank means no recorded answer, not a negative.
                columns = [i + 1 for i, c in enumerate(row.cells) if c._tc is cell._tc]
                header_values = [text(headers[i - 1]) for i in columns]
                flags = []
                if len(columns) > 1:
                    flags.append('merged_response_requires_review')
                if not context:
                    flags.append('missing_row_context')
                if any(not h for h in header_values):
                    flags.append('missing_column_header')
                records.append({
                    'record_id': f'{digest[:12]}-t{table_index}-r{row_index}-c{column_index+1}',
                    'document_sha256': digest,
                    'source_file': path.name,
                    'source_type': 'community_data',
                    'source_section': f'Table {table_index}, row {row_index}, columns {columns}',
                    'page_number': '',
                    'table_number': table_index,
                    'row_number': row_index,
                    'column_numbers': columns,
                    'context_kind': context_name,
                    'row_context': context,
                    'column_headers': header_values,
                    'response': value,
                    'review_status': 'pending',
                    'review_flags': flags,
                })
    if len(document.tables) != 3:
        warnings.append('Expected three tables; document layout needs review')
    return records, warnings


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input-dir', required=True, type=Path)
    parser.add_argument('--output-dir', required=True, type=Path)
    args = parser.parse_args()
    files = sorted(p for p in args.input_dir.iterdir()
                   if p.suffix.lower() == '.docx' and not p.name.startswith('~$'))
    if not files:
        parser.error('No Word documents found')
    args.output_dir.mkdir(parents=True, exist_ok=True)
    total = 0
    for index, path in enumerate(files, 1):
        records, warnings = extract_document(path)
        target = args.output_dir / f'community_{index:02d}_evidence.json'
        if target.exists():
            raise SystemExit('Output exists; choose a new output directory to preserve prior review.')
        with target.open('x', encoding='utf-8') as f:
            json.dump({'status': 'unreviewed_table_evidence',
                       'warnings': warnings, 'records': records}, f,
                      ensure_ascii=False, indent=2)
        total += len(records)
        flagged = sum(bool(r['review_flags']) for r in records)
        print(f'Document {index}: {len(records)} evidence cells; {flagged} flagged; {len(warnings)} layout warnings')
    print(f'Saved {total} unreviewed evidence records locally. No model calls or database uploads.')


if __name__ == '__main__':
    main()

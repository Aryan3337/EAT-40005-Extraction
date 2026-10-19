import sys
from pathlib import Path
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / 'scripts'))

import tempfile
import unittest
from pathlib import Path
from docx import Document
from community_narratives import extract_narratives

class NarrativeTests(unittest.TestCase):
    def test_question_negation_and_context_preserved(self):
        with tempfile.TemporaryDirectory() as tmp:
            d=Document(); d.add_paragraph('Section 2: Demo')
            d.add_paragraph('1. Any changes?\nAnswer: No.')
            d.add_paragraph('Unmarked response for review')
            p=Path(tmp)/'demo.docx'; d.save(p)
            out=extract_narratives(p); r=out['records'][0]
            self.assertEqual(r['question'],'1. Any changes?')
            self.assertEqual(r['response'],'No.')
            self.assertEqual(r['section_heading'],'Section 2: Demo')
            self.assertEqual(r['page_number'],'')
            self.assertEqual(len(out['unclassified_paragraphs']),1)
            self.assertEqual(r['review_status'],'pending')

    def test_missing_context_is_flagged(self):
        with tempfile.TemporaryDirectory() as tmp:
            d=Document();d.add_paragraph('Question? Answer:')
            p=Path(tmp)/'demo.docx';d.save(p)
            r=extract_narratives(p)['records'][0]
            self.assertIn('missing_response',r['review_flags'])
            self.assertIn('missing_section',r['review_flags'])

if __name__=='__main__': unittest.main()

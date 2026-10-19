import sys
from pathlib import Path
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / 'scripts'))

import csv
import json
import tempfile
import unittest
from pathlib import Path
from docx import Document
from community_narratives import extract_narratives
from community_combine import combine

class CombineTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.docs = self.root / 'docs'; self.docs.mkdir()
        self.evidence = self.root / 'evidence'; self.evidence.mkdir()
        doc = Document()
        doc.add_paragraph('Section 1: Fictional test')
        doc.add_paragraph('Any changes?\nAnswer: No.')
        doc.save(self.docs / 'fictional.docx')
        self.payload = extract_narratives(self.docs / 'fictional.docx')
        self.json_path = self.evidence / 'community_01_narratives.json'
        self.json_path.write_text(json.dumps(self.payload))
        self.tables = self.root / 'table.csv'
        fields = ['subject','predicate','object','sentence_ref','source_section',
                  'confidence','paper_title','source_type','source_file','page_number',
                  'record_id','review_status','verification_scope','time_period']
        with self.tables.open('w', newline='') as f:
            w = csv.DictWriter(f, fieldnames=fields); w.writeheader()
            w.writerow(dict(subject='Demo',predicate='REPORTED_RESPONSE',object='DemoCrop',
                            source_type='community_data',record_id='fictional-table',
                            review_status='candidate'))
        self.out = self.root / 'combined.csv'

    def tearDown(self): self.tmp.cleanup()

    def run_conversion(self):
        return combine(self.tables,self.evidence,self.docs,self.out)

    def test_combines_without_losing_no_or_question(self):
        self.assertEqual(self.run_conversion(),(1,1))
        with self.out.open(newline='') as f: rows=list(csv.DictReader(f))
        self.assertEqual(len(rows),2)
        self.assertEqual(rows[1]['object'],'No.')
        self.assertEqual(rows[1]['question'],'Any changes?')
        self.assertEqual(rows[1]['review_status'],'candidate')
        self.assertEqual(rows[1]['page_number'],'')

    def test_changed_evidence_is_rejected_before_writing(self):
        self.payload['records'][0]['response']='Altered answer'
        self.json_path.write_text(json.dumps(self.payload))
        with self.assertRaisesRegex(ValueError,'differs'):
            self.run_conversion()
        self.assertFalse(self.out.exists())

    def test_existing_output_is_preserved(self):
        self.out.write_text('Existing reviewed file')
        with self.assertRaises(FileExistsError): self.run_conversion()
        self.assertEqual(self.out.read_text(),'Existing reviewed file')

if __name__=='__main__': unittest.main()

import sys
from pathlib import Path
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / 'scripts'))

import tempfile
import unittest
from pathlib import Path
from docx import Document
from community_tables import extract_document


class ReaderTests(unittest.TestCase):
    def test_time_columns_and_context_are_separate(self):
        with tempfile.TemporaryDirectory() as tmp:
            doc = Document()
            t = doc.add_table(rows=2, cols=5)
            t.cell(0, 0).merge(t.cell(0, 1)).text = 'Category'
            for col, value in enumerate(['Past', 'Recent', 'Now'], 2):
                t.cell(0, col).text = value
            t.cell(1, 0).merge(t.cell(1, 1)).text = 'Fictional resources'
            t.cell(1, 2).text = 'DemoPast'
            t.cell(1, 4).text = 'DemoNow'
            p = Path(tmp) / 'fictional.docx'
            doc.save(p)
            rows, warnings = extract_document(p)
            self.assertEqual(len(rows), 2)
            self.assertEqual([r['column_headers'] for r in rows], [['Past'], ['Now']])
            self.assertTrue(all(r['row_context'] == 'Fictional resources' for r in rows))
            self.assertTrue(all(r['review_status'] == 'pending' for r in rows))
            self.assertTrue(all(r['page_number'] == '' for r in rows))
            self.assertTrue(warnings)  # Fixture deliberately only has one table.

    def test_merged_responses_require_review_and_are_not_duplicated(self):
        with tempfile.TemporaryDirectory() as tmp:
            doc = Document()
            t = doc.add_table(rows=2, cols=5)
            for i in range(5):
                t.cell(0, i).text = f'Header {i}'
            t.cell(1, 0).text = 'DemoContext'
            t.cell(1, 2).merge(t.cell(1, 3)).text = 'Demo response'
            p = Path(tmp) / 'fictional.docx'
            doc.save(p)
            rows, _ = extract_document(p)
            self.assertEqual(len(rows), 1)
            self.assertEqual(rows[0]['column_numbers'], [3, 4])
            self.assertIn('merged_response_requires_review', rows[0]['review_flags'])


if __name__ == '__main__':
    unittest.main()
